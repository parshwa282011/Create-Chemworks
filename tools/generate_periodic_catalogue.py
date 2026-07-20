#!/usr/bin/env python3
"""Generate the hardcoded 558-element table and designer catalogues."""

from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
JAVA = ROOT / "src/main/java/com/parswha/createchemworks/client/Elements/ElementsData.java"
CSV = ROOT / "designs/chemworks_elements_1_558.csv"
XLSX = ROOT / "designs/chemworks_elements_1_558.xlsx"

START_PERIOD = {"S": 1, "P": 2, "D": 4, "F": 6, "H": 8, "M": 10, "A": 12}
CAPACITY = {"S": 2, "P": 6, "D": 10, "F": 14, "H": 18, "M": 22, "A": 26}
COLUMN_START = {"S": 0, "A": 0, "M": 0, "H": 0, "F": 0, "D": 3, "P": 13}
PERIOD_ORDER = ("S", "A", "M", "H", "F", "D", "P")
COLORS = {"S": "2F6BFF", "P": "2DBE60", "D": "FF8C2A", "F": "A855F7",
          "H": "00B8D9", "M": "E23D55", "A": "D4A017"}
ACTIVE_WORLDS = ("Earth", "Moon", "Mercury", "Venus", "Mars")
PERIOD_ROOT = {8: "Octava", 9: "Ennea", 10: "Decara", 11: "Undeca",
               12: "Dodeca", 13: "Trideca"}
BLOCK_ROOTS = {
    "S": ("sol", "alk", "ion", "polar"),
    "P": ("aer", "cryst", "lumin", "react", "plasm", "vapor"),
    "D": ("ferro", "catal", "conduct", "chrom", "kinet", "allotrop"),
    "F": ("fiss", "radio", "actin", "nucleo", "decay", "fusion"),
    "H": ("harmon", "reson", "gyro", "sonor", "rotar", "phono"),
    "M": ("meta", "phase", "quant", "flux", "shift", "isomer"),
    "A": ("anom", "chrono", "gravi", "paradox", "axi", "void"),
}
SPECIAL = {
    119: "Ancientium", 120: "Catalyrium", 121: "Resonarium", 122: "Voltrium",
    123: "Synthium", 124: "Adaptium", 125: "Unifornium", 126: "Terralium",
    127: "Lunarium", 128: "Levitium", 129: "Pyronium", 130: "Aresium",
    131: "Harmonium", 132: "Gyrium", 133: "Sonorium", 134: "Fluxium",
    135: "Quantium", 136: "Metastium", 137: "Phasium", 138: "Anomalium",
}
SPECIAL_SYMBOLS = {119: "An"}
HEADERS = ("atomic_number", "symbol", "element_name", "period", "electron_block",
           "column", "production_rule", "celestial_source", "designer_status", "notes")


def real_elements() -> list[dict[str, object]]:
    text = JAVA.read_text(encoding="utf-8")
    pattern = re.compile(
        r'(?:m|s|e)\((\d+),\s*"([^"]+)",\s*"([^"]+)",\s*(\d+),\s*(\d+)'
        r'(?:,\s*ElectronBlock\.([A-Z]))?\)'
    )
    found = []
    for match in pattern.finditer(text):
        number, symbol, name, period, column, block = match.groups()
        number = int(number)
        if number > 118:
            continue
        if block is None:
            block = "F"
        found.append(row(number, symbol, name, int(period), block, int(column), True))
    unique = {int(value["atomic_number"]): value for value in found}
    if len(unique) != 118:
        raise RuntimeError(f"Could not recover 118 real elements from ElementsData.java: {len(unique)}")
    ordered = [unique[number] for number in range(1, 119)]

    # Rebuild columns from the element order instead of repeatedly translating
    # the columns found in the previously generated Java file. This makes the
    # generator idempotent and prevents later P/D elements sharing a cell.
    positions: dict[tuple[int, str], int] = {}
    for value in ordered:
        period = int(value["period"])
        block = str(value["electron_block"]).upper()
        key = (period, block)
        position = positions.get(key, 0)
        value["column"] = COLUMN_START[block] + position
        positions[key] = position + 1
    ordered[1]["column"] = 18  # Helium occupies the final P-side column.
    return ordered


def fictional_symbols(occupied: set[str]) -> list[str]:
    """Return compact, unique symbols that remain readable in the GUI cells."""
    symbols = []
    for first in "ABCDEFGHIJKLMNOPQRSTUVWXYZ":
        for second in "abcdefghijklmnopqrstuvwxyz":
            symbol = first + second
            if symbol not in occupied:
                symbols.append(symbol)
    return symbols


def fictional_name(number: int, period: int, block: str, position: int) -> str:
    if number in SPECIAL:
        return SPECIAL[number]
    root = PERIOD_ROOT[period]
    family = BLOCK_ROOTS[block][position % len(BLOCK_ROOTS[block])]
    cycle = position // len(BLOCK_ROOTS[block])
    syllable = "" if cycle == 0 else chr(ord("a") + cycle)
    ending = "um" if family.endswith("i") else "ium"
    return f"{root}{family}{syllable}{ending}"


def row(number: int, symbol: str, name: str, period: int, block: str,
        column: int, real: bool = False) -> dict[str, object]:
    reaction_only = period >= 10
    if real:
        source = "Real-world element"
        rule = "Existing real element"
    elif reaction_only:
        source = "Particle Accelerator"
        rule = "Particle Reaction only"
    else:
        source = ACTIVE_WORLDS[(number + column) % len(ACTIVE_WORLDS)]
        rule = "Natural deposit or processing"
    return {
        "atomic_number": number, "symbol": symbol, "element_name": name,
        "period": period, "electron_block": block.lower(), "column": column,
        "production_rule": rule, "celestial_source": source,
        "designer_status": "Not started",
        "notes": "Real element" if real else "Pure fictional element; not an alloy",
    }


def build() -> list[dict[str, object]]:
    rows = real_elements()
    occupied_symbols = {str(value["symbol"]) for value in rows} | set(SPECIAL_SYMBOLS.values())
    symbols = iter(fictional_symbols(occupied_symbols))
    number = 119
    for period in range(8, 14):
        for block in PERIOD_ORDER:
            if period < START_PERIOD[block]:
                continue
            for position in range(CAPACITY[block]):
                symbol = SPECIAL_SYMBOLS[number] if number in SPECIAL_SYMBOLS else next(symbols)
                rows.append(row(
                    number, symbol,
                    fictional_name(number, period, block, position),
                    period, block, COLUMN_START[block] + position,
                ))
                number += 1
    assert number == 559 and len(rows) == 558
    assert len({value["symbol"] for value in rows}) == 558
    assert len({value["element_name"] for value in rows}) == 558
    cells = [(value["period"], value["electron_block"], value["column"]) for value in rows]
    assert len(cells) == len(set(cells)), "Two elements occupy the same block cell"
    return rows


def java_source(rows: list[dict[str, object]]) -> str:
    entries = []
    for value in rows:
        entries.append(
            f'        e({value["atomic_number"]}, "{value["symbol"]}", '
            f'"{value["element_name"]}", {value["period"]}, {value["column"]}, '
            f'ElectronBlock.{str(value["electron_block"]).upper()})'
        )
    joined = ",\n".join(entries)
    return f"""package com.parswha.createchemworks.client.Elements;

import java.util.List;

public final class ElementsData {{
    public static final int PERIODS = 13;
    public static final int MAIN_COLUMNS = 19;

    // Intentionally hardcoded while the extended chemistry design is in flux.
    private static final List<Element> ELEMENTS = List.of(
{joined}
    );

    private ElementsData() {{}}

    public static List<Element> all() {{ return ELEMENTS; }}

    public static List<Element> mainTableElements() {{
        return ELEMENTS.stream().filter(Element::belongsOnMainTable).toList();
    }}

    public static List<Element> elementsIn(ElectronBlock block) {{
        return ELEMENTS.stream().filter(element -> element.block() == block).toList();
    }}

    private static Element e(int number, String symbol, String name, int period,
                             int column, ElectronBlock block) {{
        return new Element(number, symbol, name, period, column, block);
    }}
}}
"""


def write_outputs(rows: list[dict[str, object]]) -> None:
    JAVA.write_text(java_source(rows), encoding="utf-8")
    CSV.parent.mkdir(parents=True, exist_ok=True)
    with CSV.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows(rows)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Elements 1-558"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:J{len(rows) + 1}"
    for column, header in enumerate(HEADERS, 1):
        cell = sheet.cell(1, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="20242B")
    for row_index, value in enumerate(rows, 2):
        for column, header in enumerate(HEADERS, 1):
            sheet.cell(row_index, column, value[header])
        color = COLORS[str(value["electron_block"]).upper()]
        sheet.cell(row_index, 5).fill = PatternFill("solid", fgColor=color)
        sheet.cell(row_index, 5).font = Font(bold=True, color="FFFFFF")
    widths = (14, 10, 29, 10, 16, 10, 28, 24, 18, 40)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    guide = workbook.create_sheet("Rules")
    rules = [
        ("Rule", "Value"),
        ("Total elements", "558"),
        ("s-block", "Periods 1-13; 26 elements"),
        ("p-block", "Periods 2-13; 72 elements"),
        ("d-block", "Periods 4-13; 100 elements"),
        ("f-block", "Periods 6-13; 112 elements"),
        ("h-block", "Periods 8-13; 108 elements"),
        ("m-block", "Periods 10-13; 88 elements"),
        ("a-block", "Periods 12-13; 52 elements"),
        ("Production lock", "Every element in periods 10-13 is Particle Reaction only"),
    ]
    for value in rules:
        guide.append(value)
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 72
    workbook.save(XLSX)


def main() -> None:
    rows = build()
    write_outputs(rows)
    print(f"Generated {len(rows)} hardcoded elements")
    print(JAVA)
    print(CSV)
    print(XLSX)


if __name__ == "__main__":
    main()
